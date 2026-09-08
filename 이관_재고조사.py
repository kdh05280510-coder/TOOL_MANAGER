# D:\tool_manager\이관_재고조사.py
# 사용:
#   1) 공구등록 프로그램, exe, DB Browser 종료
#   2) 엑셀을 D:\tool_manager 에 두기
#   3) cd /d D:\tool_manager
#      python 이관_재고조사.py

import re
import sys
import traceback
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook
from database.db import get_connection, init_db

EXCEL_NAMES = [
    "공구관리대장_성진세미텍_재고조사.xlsm",
    "공구관리대장_성진세미텍.xlsm",
    "공구관리대장_성진세미텍ver2.xlsm",
]
SHEET_NAME = "재고조사"
IS_GRADE_B = 0  # 0=A급, 1=B급


def s(v):
    if v is None:
        return ""
    return str(v).strip()


def to_float(v):
    text = s(v)
    if not text or text == "None":
        return None
    text = (
        text.replace("생크지름:", "")
        .replace("생크직경:", "")
        .replace("전체길이:", "")
        .replace("온길이:", "")
    )
    try:
        return float(text)
    except ValueError:
        m = re.search(r"([0-9.]+)", text)
        return float(m.group(1)) if m else None


def parse_dl(tool_name):
    text = s(tool_name)
    d = l = None
    m = re.search(r"[Dd]\s*([0-9.]+)", text)
    if m:
        d = float(m.group(1))
    m = re.search(r"[Ll]\s*([0-9.]+)", text)
    if m:
        l = float(m.group(1))
    return d, l


def get_or_create_maker(cur, name):
    name = s(name)
    if not name or name in ("-", "?", "None"):
        return None
    cur.execute("SELECT id FROM makers WHERE name = ?", (name,))
    row = cur.fetchone()
    if row:
        return row["id"]
    cur.execute("INSERT INTO makers (name, is_active) VALUES (?, 1)", (name,))
    return cur.lastrowid


def get_or_create_category(cur, main_name, sub_code):
    main_name = s(main_name)
    sub_code = s(sub_code)
    if not main_name or not sub_code or main_name == "None" or sub_code == "None":
        return None
    cur.execute(
        "SELECT id FROM categories WHERE main_name=? AND sub_code=?",
        (main_name, sub_code),
    )
    row = cur.fetchone()
    if row:
        return row["id"]
    main_code = main_name.split("(")[0].strip() if "(" in main_name else main_name
    cur.execute(
        "INSERT INTO categories (main_code, main_name, sub_code) VALUES (?, ?, ?)",
        (main_code, main_name, sub_code),
    )
    print("  분류 추가:", main_name, "/", sub_code)
    return cur.lastrowid


def find_excel():
    for name in EXCEL_NAMES:
        p = ROOT / name
        if p.exists():
            return p
    return None


def main():
    try:
        init_db()
        path = find_excel()
        if not path:
            print("엑셀 파일이 없습니다. 아래 이름으로 D:\\tool_manager 에 두세요.")
            for name in EXCEL_NAMES:
                print(" -", ROOT / name)
            return

        print("엑셀:", path)
        wb = load_workbook(path, data_only=True)
        print("시트 목록:", wb.sheetnames)
        if SHEET_NAME not in wb.sheetnames:
            print("시트 없음:", SHEET_NAME)
            return

        ws = wb[SHEET_NAME]
        print("처리 중:", SHEET_NAME)

        conn = get_connection()
        cur = conn.cursor()
        tools_ok = inv_ok = skip = 0

        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row:
                skip += 1
                continue

            main_name = row[1] if len(row) > 1 else None
            sub_code = row[2] if len(row) > 2 else None
            maker_name = row[3] if len(row) > 3 else None
            tool_name = row[4] if len(row) > 4 else None
            sub_name = row[5] if len(row) > 5 else None
            tool_code = row[6] if len(row) > 6 else None
            barcode = s(row[7] if len(row) > 7 else None)
            shank = row[9] if len(row) > 9 else None
            total = row[10] if len(row) > 10 else None

            if not barcode or barcode == "None":
                skip += 1
                continue

            cur.execute("SELECT id FROM inventory WHERE barcode = ?", (barcode,))
            if cur.fetchone():
                print("  이미 있음 스킵:", barcode)
                skip += 1
                continue

            cat_id = get_or_create_category(cur, main_name, sub_code)
            if not cat_id:
                print("  분류 없음 스킵:", main_name, sub_code, barcode)
                skip += 1
                continue

            maker_id = get_or_create_maker(cur, maker_name)
            diameter, length = parse_dl(tool_name)

            cur.execute(
                """
                INSERT INTO tools (
                    category_id, maker_id, tool_code, tool_name,
                    diameter, length, shank_dia, total_length
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    cat_id,
                    maker_id,
                    s(tool_code) or None,
                    s(tool_name) or None,
                    diameter,
                    length,
                    to_float(shank),
                    to_float(total),
                ),
            )
            tool_id = cur.lastrowid
            tools_ok += 1

            cur.execute(
                """
                INSERT INTO inventory (
                    tool_id, barcode, sub_name, quantity, status, is_grade_b
                ) VALUES (?, ?, ?, 1, ?, ?)
                """,
                (
                    tool_id,
                    barcode,
                    s(sub_name) or None,
                    "B급" if IS_GRADE_B else "정상",
                    1 if IS_GRADE_B else 0,
                ),
            )
            inv_ok += 1
            print("  추가:", barcode, s(tool_name), s(main_name), s(sub_code))

        conn.commit()
        conn.close()
        print("이관 완료")
        print("tools 추가:", tools_ok)
        print("inventory 추가:", inv_ok)
        print("스킵:", skip)
    except Exception:
        print("----- 오류 -----")
        traceback.print_exc()
        print("----------------")
        print("프로그램/DB Browser를 끄고 다시 실행하세요.")
        print("database is locked 이면 tools.db 를 다른 프로그램이 연 상태입니다.")


if __name__ == "__main__":
    main()
